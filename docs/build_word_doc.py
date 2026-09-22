"""Build a Word document with all F365 app page screenshots."""
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SCREENSHOTS_DIR = os.path.join(os.path.dirname(__file__), "..", "page-screenshots")
OUTPUT = os.path.join(os.path.dirname(__file__), "F365-App-All-Pages-Screenshots.docx")


def title_from_filename(name: str) -> str:
    base = os.path.splitext(name)[0]
    return base.replace("-", " ").replace("_", " ").title()


def main() -> None:
    files = sorted(f for f in os.listdir(SCREENSHOTS_DIR) if f.lower().endswith(".png"))

    doc = Document()

    # Cover
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("F365 — All App Pages")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0xC2, 0x18, 0x5B)

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = s.add_run("Visual reference of every screen in the app")
    sr.font.size = Pt(13)
    sr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run(f"Total screens: {len(files)}")
    mr.italic = True
    mr.font.size = Pt(11)

    doc.add_page_break()

    # Table of contents
    toc = doc.add_paragraph()
    tr = toc.add_run("Contents")
    tr.bold = True
    tr.font.size = Pt(18)
    for f in files:
        doc.add_paragraph(title_from_filename(f), style="List Number")

    doc.add_page_break()

    # Pages
    for f in files:
        path = os.path.join(SCREENSHOTS_DIR, f)
        h = doc.add_paragraph()
        hr = h.add_run(title_from_filename(f))
        hr.bold = True
        hr.font.size = Pt(20)
        hr.font.color.rgb = RGBColor(0xC2, 0x18, 0x5B)

        sub = doc.add_paragraph()
        sb = sub.add_run(f)
        sb.italic = True
        sb.font.size = Pt(9)
        sb.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(path, height=Inches(7.0))

        doc.add_page_break()

    doc.save(OUTPUT)
    print(f"Saved: {OUTPUT}")


if __name__ == "__main__":
    main()
